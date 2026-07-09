import os
import re
import fitz          # PDF — pymupdf
import docx          # Word — python-docx
import openpyxl      # Excel

def clean_text(text: str) -> str:
    text = re.sub(r'[^\x20-\x7E\n\r\t]', ' ', text)
    text = re.sub(r' {2,}', ' ', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()

def extract_from_pdf(file_path: str) -> str:
    all_text = []
    with fitz.open(file_path) as doc:
        total_pages = len(doc)
        for page_num, page in enumerate(doc, start=1):
            text = page.get_text("text")
            if not text.strip() or len(text) < 50:
                blocks = page.get_text("blocks")
                text = "\n".join(
                    b[4] for b in blocks if isinstance(b[4], str)
                )
            text = clean_text(text)
            if text.strip():
                all_text.append(f"--- Page {page_num} ---\n{text}")
    full_text = "\n".join(all_text)
    print(f"PDF: Extracted {len(full_text)} chars from {total_pages} pages")
    return full_text

def extract_from_docx(file_path: str) -> str:
    doc = docx.Document(file_path)
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            # Preserve heading structure
            if para.style.name.startswith("Heading"):
                parts.append(f"\n=== {para.text.strip()} ===\n")
            else:
                parts.append(para.text.strip())
    # Also extract tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip() for cell in row.cells
                if cell.text.strip()
            )
            if row_text:
                parts.append(row_text)
    full_text = clean_text("\n".join(parts))
    print(f"DOCX: Extracted {len(full_text)} chars")
    return full_text

def extract_from_xlsx(file_path: str) -> str:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    all_text = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        all_text.append(f"\n=== Sheet: {sheet_name} ===\n")
        for row in ws.iter_rows(values_only=True):
            row_text = " | ".join(
                str(cell) for cell in row if cell is not None
            )
            if row_text.strip():
                all_text.append(row_text)
    full_text = clean_text("\n".join(all_text))
    print(f"XLSX: Extracted {len(full_text)} chars from {len(wb.sheetnames)} sheets")
    return full_text

def extract_from_txt(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read()
    full_text = clean_text(text)
    print(f"TXT: Extracted {len(full_text)} chars")
    return full_text

# Supported extensions
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".txt"}

def extract_text(file_path: str) -> str:
    """Route to correct parser based on file extension."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return extract_from_pdf(file_path)
    elif ext == ".docx":
        return extract_from_docx(file_path)
    elif ext == ".xlsx":
        return extract_from_xlsx(file_path)
    elif ext == ".txt":
        return extract_from_txt(file_path)
    else:
        raise ValueError(
            f"Unsupported file type '{ext}'. "
            f"Supported: {SUPPORTED_EXTENSIONS}"
        )